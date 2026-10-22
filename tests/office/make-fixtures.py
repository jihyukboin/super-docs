"""Regenerate independent fixtures using openpyxl and python-pptx (test tooling only)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches
root = Path('tests/office/fixtures')
w = Workbook(); s = w.active; s.title = '매출'
s.append(['제품', '수량', '단가', '합계'])
s.append(['사과', 3, 1200, '=B2*C2'])
s.append(['배', 2, 2500, '=B3*C3'])
s['A1'].font = Font(bold=True, color='FFFFFF'); s['A1'].fill = PatternFill('solid', fgColor='336699')
s.merge_cells('A5:C5'); s['A5'] = '합계'; s['D5'] = '=SUM(D2:D3)'
s.column_dimensions['A'].width = 24
chart = BarChart(); chart.add_data(Reference(s, min_col=2, min_row=1, max_row=3), titles_from_data=True); s.add_chart(chart, 'F2')
w.create_sheet('메모')['A1'] = '보존할 두 번째 시트'
w.save(root/'sample.xlsx')
p = Presentation(); slide = p.slides.add_slide(p.slide_layouts[5]); slide.shapes.title.text = 'Quarterly Report'
box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)); box.text = 'Editable slide text'
table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(5), Inches(1)).table
table.cell(0,0).text = '제품'; table.cell(0,1).text = '금액'; table.cell(1,0).text = '사과'; table.cell(1,1).text = '3600'
slide.notes_slide.notes_text_frame.text = 'Preserve speaker notes'
second = p.slides.add_slide(p.slide_layouts[1]); second.shapes.title.text = 'Second slide'
p.save(root/'sample.pptx')
